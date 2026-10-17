#!/usr/bin/env python3
"""
reader_generator.py
Generates standalone, in-browser full-text HTML reader pages for all 60 Value Physics deliverables (VPE-01 to VPE-60).
"""

import os
import sys
import json
import html
import zipfile
import re
import xml.etree.ElementTree as ET

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DELIV_DIR = os.path.join(BASE_DIR, "deliverables")
READ_DIR = os.path.join(BASE_DIR, "read")

READER_CSS = """
:root {
  --bg-primary: #0a0e17;
  --bg-secondary: #121824;
  --bg-card: #111827;
  --bg-card-subtle: #1a2234;
  --border-color: rgba(255, 255, 255, 0.08);
  --border-active: rgba(0, 212, 255, 0.4);
  --text-main: #f1f5f9;
  --text-muted: #94a3b8;
  --text-dim: #64748b;
  --accent-cyan: #00d4ff;
  --accent-gold: #f59e0b;
  --accent-green: #10b981;
  --accent-purple: #a855f7;
  --accent-red: #ef4444;
  --font-mono: 'JetBrains Mono', 'Fira Code', monospace;
  --font-sans: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
}

* {
  box-sizing: border-box;
  margin: 0;
  padding: 0;
}

body {
  background: var(--bg-primary);
  color: var(--text-main);
  font-family: var(--font-sans);
  line-height: 1.7;
  min-height: 100vh;
  display: flex;
  flex-direction: column;
}

/* Sticky Navigation Header */
.reader-nav {
  position: sticky;
  top: 0;
  z-index: 100;
  background: rgba(10, 14, 23, 0.92);
  backdrop-filter: blur(12px);
  border-bottom: 1px solid var(--border-color);
  padding: 0.75rem 1.5rem;
  display: flex;
  align-items: center;
  justify-content: space-between;
  gap: 1rem;
}

.reader-nav-left {
  display: flex;
  align-items: center;
  gap: 0.85rem;
  flex-wrap: wrap;
}

.nav-back-link {
  display: inline-flex;
  align-items: center;
  gap: 0.4rem;
  color: var(--text-muted);
  text-decoration: none;
  font-size: 0.85rem;
  font-weight: 500;
  padding: 0.35rem 0.75rem;
  border-radius: 6px;
  background: rgba(255, 255, 255, 0.04);
  border: 1px solid var(--border-color);
  transition: all 0.2s ease;
}

.nav-back-link:hover {
  color: var(--accent-cyan);
  border-color: var(--accent-cyan);
  background: rgba(0, 212, 255, 0.08);
}

.doc-badge-id {
  font-family: var(--font-mono);
  font-weight: 700;
  font-size: 0.85rem;
  background: linear-gradient(135deg, rgba(0, 212, 255, 0.2), rgba(37, 99, 235, 0.2));
  color: var(--accent-cyan);
  border: 1px solid rgba(0, 212, 255, 0.3);
  padding: 0.25rem 0.65rem;
  border-radius: 5px;
}

.doc-badge-cat {
  font-size: 0.8rem;
  color: var(--accent-gold);
  background: rgba(245, 158, 11, 0.1);
  border: 1px solid rgba(245, 158, 11, 0.25);
  padding: 0.2rem 0.6rem;
  border-radius: 4px;
}

.doc-badge-status {
  font-size: 0.75rem;
  color: var(--accent-green);
  background: rgba(16, 185, 129, 0.1);
  border: 1px solid rgba(16, 185, 129, 0.25);
  padding: 0.2rem 0.55rem;
  border-radius: 4px;
}

.reader-nav-right {
  display: flex;
  align-items: center;
  gap: 0.6rem;
}

.nav-arrow-btn {
  display: inline-flex;
  align-items: center;
  gap: 0.35rem;
  padding: 0.35rem 0.75rem;
  font-size: 0.8rem;
  font-weight: 500;
  text-decoration: none;
  color: var(--text-muted);
  border-radius: 6px;
  background: rgba(255, 255, 255, 0.04);
  border: 1px solid var(--border-color);
  transition: all 0.2s ease;
}

.nav-arrow-btn:hover:not(.disabled) {
  color: var(--text-main);
  border-color: rgba(255, 255, 255, 0.2);
  background: rgba(255, 255, 255, 0.08);
}

.nav-arrow-btn.disabled {
  opacity: 0.35;
  cursor: not-allowed;
  pointer-events: none;
}

.nav-download-btn {
  display: inline-flex;
  align-items: center;
  gap: 0.45rem;
  padding: 0.38rem 0.9rem;
  font-size: 0.82rem;
  font-weight: 600;
  text-decoration: none;
  color: #0a0e17;
  background: linear-gradient(135deg, #00d4ff, #38bdf8);
  border-radius: 6px;
  transition: all 0.2s ease;
}

.nav-download-btn:hover {
  box-shadow: 0 0 12px rgba(0, 212, 255, 0.4);
  transform: translateY(-1px);
}

/* Main Container */
.reader-container {
  max-width: 960px;
  width: 100%;
  margin: 2rem auto;
  padding: 0 1.5rem;
  flex: 1;
}

/* Article Card */
.article-card {
  background: var(--bg-card);
  border: 1px solid var(--border-color);
  border-radius: 12px;
  padding: 3rem 3.5rem;
  box-shadow: 0 10px 30px rgba(0, 0, 0, 0.35);
}

@media (max-width: 768px) {
  .article-card {
    padding: 1.5rem;
  }
  .reader-nav {
    padding: 0.75rem 1rem;
  }
}

/* Article Header */
.article-header {
  border-bottom: 1px solid var(--border-color);
  padding-bottom: 2rem;
  margin-bottom: 2.5rem;
}

.article-header-meta {
  display: flex;
  align-items: center;
  gap: 0.75rem;
  margin-bottom: 1rem;
  flex-wrap: wrap;
}

.article-title {
  font-size: 2rem;
  font-weight: 800;
  line-height: 1.3;
  color: var(--text-main);
  margin-bottom: 1.25rem;
  letter-spacing: -0.02em;
}

.article-author-box {
  background: rgba(255, 255, 255, 0.02);
  border: 1px solid rgba(255, 255, 255, 0.06);
  border-radius: 8px;
  padding: 1rem 1.25rem;
  margin-top: 1.25rem;
}

.article-author-name {
  font-weight: 600;
  color: var(--accent-cyan);
  font-size: 0.95rem;
}

.article-author-inst {
  color: var(--text-muted);
  font-size: 0.85rem;
  margin-top: 0.2rem;
}

.article-scope-box {
  margin-top: 1rem;
  padding-top: 1rem;
  border-top: 1px dashed var(--border-color);
  font-size: 0.9rem;
  color: var(--text-muted);
  line-height: 1.6;
}

.article-data-pills {
  display: flex;
  flex-wrap: wrap;
  gap: 0.4rem;
  margin-top: 0.85rem;
}

.data-pill {
  font-size: 0.75rem;
  background: rgba(0, 212, 255, 0.06);
  border: 1px solid rgba(0, 212, 255, 0.2);
  color: var(--accent-cyan);
  padding: 0.2rem 0.55rem;
  border-radius: 4px;
}

/* Document Typography & Elements */
.doc-content {
  font-size: 1.05rem;
  line-height: 1.8;
  color: #e2e8f0;
}

.doc-h1, .doc-h2 {
  font-size: 1.55rem;
  font-weight: 700;
  color: #ffffff;
  margin-top: 2.5rem;
  margin-bottom: 1rem;
  padding-bottom: 0.5rem;
  border-bottom: 1px solid rgba(255, 255, 255, 0.1);
  letter-spacing: -0.01em;
}

.doc-h3 {
  font-size: 1.25rem;
  font-weight: 600;
  color: var(--accent-cyan);
  margin-top: 1.8rem;
  margin-bottom: 0.75rem;
}

.doc-h4 {
  font-size: 1.08rem;
  font-weight: 600;
  color: var(--accent-gold);
  margin-top: 1.4rem;
  margin-bottom: 0.5rem;
}

.doc-p {
  margin-bottom: 1.25rem;
  color: #e2e8f0;
  text-align: justify;
}

.doc-p strong {
  color: #ffffff;
  font-weight: 600;
}

.doc-p em {
  color: #cbd5e1;
}

.doc-ul, .doc-ol {
  margin: 1rem 0 1.5rem 1.8rem;
  color: #e2e8f0;
}

.doc-li {
  margin-bottom: 0.5rem;
  line-height: 1.7;
}

.doc-li strong {
  color: #ffffff;
}

/* Table Styling */
.table-container {
  width: 100%;
  overflow-x: auto;
  margin: 1.8rem 0;
  border-radius: 8px;
  border: 1px solid var(--border-color);
  background: var(--bg-secondary);
}

.doc-table {
  width: 100%;
  border-collapse: collapse;
  font-size: 0.9rem;
  text-align: left;
}

.doc-table th {
  background: rgba(255, 255, 255, 0.05);
  color: var(--accent-cyan);
  font-weight: 600;
  padding: 0.75rem 1rem;
  border-bottom: 1px solid var(--border-color);
  white-space: nowrap;
}

.doc-table td {
  padding: 0.75rem 1rem;
  border-bottom: 1px solid rgba(255, 255, 255, 0.04);
  color: #cbd5e1;
  vertical-align: top;
}

.doc-table tr:hover {
  background: rgba(255, 255, 255, 0.02);
}

/* Code & Pre */
.code-container {
  position: relative;
  margin: 1.5rem 0;
}

.code-header {
  display: flex;
  justify-content: space-between;
  align-items: center;
  background: #0d131f;
  border: 1px solid var(--border-color);
  border-bottom: none;
  border-top-left-radius: 8px;
  border-top-right-radius: 8px;
  padding: 0.5rem 1rem;
  font-family: var(--font-mono);
  font-size: 0.8rem;
  color: var(--text-dim);
}

.copy-btn {
  background: rgba(255, 255, 255, 0.06);
  border: 1px solid var(--border-color);
  color: var(--text-muted);
  border-radius: 4px;
  padding: 0.2rem 0.55rem;
  font-size: 0.75rem;
  cursor: pointer;
  transition: all 0.2s;
}

.copy-btn:hover {
  color: var(--accent-cyan);
  border-color: var(--accent-cyan);
}

.code-block {
  background: #0d131f;
  border: 1px solid var(--border-color);
  border-bottom-left-radius: 8px;
  border-bottom-right-radius: 8px;
  padding: 1.25rem;
  overflow-x: auto;
  font-family: var(--font-mono);
  font-size: 0.88rem;
  line-height: 1.6;
  color: #e2e8f0;
}

/* PDF Viewer Frame */
.pdf-viewer-wrap {
  margin: 1.5rem 0;
  border-radius: 8px;
  overflow: hidden;
  border: 1px solid var(--border-color);
}

.pdf-toolbar {
  background: var(--bg-secondary);
  border-bottom: 1px solid var(--border-color);
  padding: 0.65rem 1rem;
  display: flex;
  justify-content: space-between;
  align-items: center;
  font-size: 0.85rem;
  color: var(--text-muted);
}

.pdf-frame {
  width: 100%;
  height: 85vh;
  border: none;
  background: #ffffff;
}

/* Slide Cards for PPTX */
.slide-card {
  background: var(--bg-secondary);
  border: 1px solid var(--border-color);
  border-radius: 8px;
  padding: 1.5rem 2rem;
  margin-bottom: 1.5rem;
}

.slide-num {
  font-family: var(--font-mono);
  font-size: 0.75rem;
  color: var(--accent-gold);
  text-transform: uppercase;
  margin-bottom: 0.5rem;
}

.slide-title {
  font-size: 1.25rem;
  font-weight: 700;
  color: var(--text-main);
  margin-bottom: 1rem;
}

/* Image container */
.image-container {
  text-align: center;
  margin: 2rem 0;
  background: var(--bg-secondary);
  border: 1px solid var(--border-color);
  border-radius: 8px;
  padding: 1.5rem;
}

.doc-image {
  max-width: 100%;
  height: auto;
  border-radius: 6px;
  border: 1px solid rgba(255, 255, 255, 0.08);
}

.image-caption {
  margin-top: 0.85rem;
  font-size: 0.85rem;
  color: var(--text-muted);
}

/* Archive file list */
.archive-list {
  background: var(--bg-secondary);
  border: 1px solid var(--border-color);
  border-radius: 8px;
  padding: 1.25rem;
  margin: 1.5rem 0;
}

/* Reader Footer Navigation */
.reader-footer {
  margin-top: 3.5rem;
  padding-top: 2rem;
  border-top: 1px solid var(--border-color);
  display: flex;
  justify-content: space-between;
  gap: 1rem;
  flex-wrap: wrap;
}

.footer-nav-card {
  flex: 1;
  min-width: 250px;
  background: var(--bg-secondary);
  border: 1px solid var(--border-color);
  border-radius: 8px;
  padding: 1.25rem;
  text-decoration: none;
  transition: all 0.2s ease;
}

.footer-nav-card:hover {
  border-color: var(--accent-cyan);
  transform: translateY(-2px);
}

.footer-nav-label {
  font-size: 0.75rem;
  text-transform: uppercase;
  color: var(--text-dim);
  font-weight: 600;
  margin-bottom: 0.35rem;
}

.footer-nav-title {
  font-size: 0.95rem;
  font-weight: 600;
  color: var(--text-main);
  line-height: 1.4;
}

/* Site Footer */
.site-footer {
  text-align: center;
  padding: 2.5rem 1.5rem;
  color: var(--text-dim);
  font-size: 0.82rem;
  border-top: 1px solid var(--border-color);
  margin-top: 3rem;
  background: var(--bg-secondary);
}

.site-footer a {
  color: var(--accent-cyan);
  text-decoration: none;
}
"""

def parse_docx_to_html(fpath):
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    html_parts = []
    
    with zipfile.ZipFile(fpath) as z:
        if 'word/document.xml' not in z.namelist():
            return "<p class='doc-p'><em>Unable to read document contents (missing word/document.xml).</em></p>"
        xml_content = z.read('word/document.xml')
        
    root = ET.fromstring(xml_content)
    body = root.find('w:body', ns)
    if body is None:
        return "<p class='doc-p'><em>Empty document.</em></p>"
        
    in_list = False
    
    for elem in body:
        tag = elem.tag.split('}')[-1]
        
        if tag == 'p':
            style_el = elem.find('.//w:pStyle', ns)
            style = style_el.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '') if style_el is not None else ''
            
            runs_html = []
            for r in elem.findall('.//w:r', ns):
                r_text = ''.join(r.itertext())
                if not r_text:
                    continue
                clean_text = html.unescape(r_text)
                esc = html.escape(clean_text)
                is_bold = r.find('.//w:b', ns) is not None
                is_italic = r.find('.//w:i', ns) is not None
                if is_bold: esc = f"<strong>{esc}</strong>"
                if is_italic: esc = f"<em>{esc}</em>"
                runs_html.append(esc)
                
            p_text = ''.join(runs_html).strip()
            if not p_text:
                continue
                
            num_el = elem.find('.//w:numPr', ns)
            is_bullet = ('List' in style) or ('bullet' in style.lower()) or (num_el is not None)
            
            if is_bullet:
                if not in_list:
                    html_parts.append('<ul class="doc-ul">')
                    in_list = True
                clean_bullet_text = re.sub(r'^[•\-\*–]\s*', '', p_text)
                html_parts.append(f'<li class="doc-li">{clean_bullet_text}</li>')
                continue
            else:
                if in_list:
                    html_parts.append('</ul>')
                    in_list = False
                    
            s_lower = style.lower()
            if 'heading1' in s_lower or 'title' in s_lower:
                html_parts.append(f'<h2 class="doc-h2">{p_text}</h2>')
            elif 'heading2' in s_lower or 'subtitle' in s_lower:
                html_parts.append(f'<h3 class="doc-h3">{p_text}</h3>')
            elif 'heading3' in s_lower:
                html_parts.append(f'<h4 class="doc-h4">{p_text}</h4>')
            else:
                html_parts.append(f'<p class="doc-p">{p_text}</p>')
                
        elif tag == 'tbl':
            if in_list:
                html_parts.append('</ul>')
                in_list = False
                
            html_parts.append('<div class="table-container"><table class="doc-table">')
            for r_idx, tr in enumerate(elem.findall('w:tr', ns)):
                html_parts.append('<tr>')
                is_hdr = (r_idx == 0) or (tr.find('.//w:tblHeader', ns) is not None)
                tag_cell = 'th' if is_hdr else 'td'
                for tc in tr.findall('w:tc', ns):
                    cell_text_parts = []
                    for p in tc.findall('w:p', ns):
                        txt = ''.join(p.itertext()).strip()
                        if txt:
                            cell_text_parts.append(html.escape(txt))
                    tc_text = '<br>'.join(cell_text_parts)
                    html_parts.append(f'<{tag_cell}>{tc_text}</{tag_cell}>')
                html_parts.append('</tr>')
            html_parts.append('</table></div>')
            
    if in_list:
        html_parts.append('</ul>')
        
    return '\n'.join(html_parts)

def parse_pptx_to_html(fpath):
    slides_html = []
    with zipfile.ZipFile(fpath) as z:
        slides = [n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]
        
        def slide_key(s):
            m = re.search(r'slide(\d+)\.xml', s)
            return int(m.group(1)) if m else 0
            
        sorted_slides = sorted(slides, key=slide_key)
        for s_idx, sname in enumerate(sorted_slides, start=1):
            root = ET.fromstring(z.read(sname))
            paragraphs = []
            for p in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}p'):
                t = ''.join(elem.text for elem in p.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t') if elem.text).strip()
                if t:
                    paragraphs.append(html.escape(t))
                    
            if not paragraphs:
                continue
                
            title = paragraphs[0] if paragraphs else f"Slide {s_idx}"
            body_lines = paragraphs[1:]
            
            card_parts = [
                f'<div class="slide-card">',
                f'  <div class="slide-num">SLIDE {s_idx:02d} / {len(sorted_slides):02d}</div>',
                f'  <h3 class="slide-title">{title}</h3>',
                '  <ul class="doc-ul">'
            ]
            for line in body_lines:
                card_parts.append(f'    <li class="doc-li">{line}</li>')
            card_parts.append('  </ul>')
            card_parts.append('</div>')
            slides_html.append('\n'.join(card_parts))
            
    return '\n'.join(slides_html)

def parse_xlsx_to_html(fpath):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        html_parts = []
        for sname in wb.sheetnames:
            sheet = wb[sname]
            html_parts.append(f'<h3 class="doc-h3">Worksheet: {html.escape(sname)}</h3>')
            html_parts.append('<div class="table-container"><table class="doc-table">')
            
            rows = list(sheet.iter_rows(values_only=True))
            header_done = False
            rendered_rows = 0
            max_rows = 150
            
            for r in rows:
                if not any(c is not None for c in r):
                    continue
                clean_row = [str(c).strip() if c is not None else "" for c in r]
                # Filter trailing empty cells
                while clean_row and not clean_row[-1]:
                    clean_row.pop()
                if not clean_row:
                    continue
                    
                tag = 'th' if not header_done else 'td'
                html_parts.append('<tr>')
                for val in clean_row:
                    html_parts.append(f'<{tag}>{html.escape(val)}</{tag}>')
                html_parts.append('</tr>')
                header_done = True
                rendered_rows += 1
                if rendered_rows >= max_rows:
                    html_parts.append(f'<tr><td colspan="{len(clean_row)}" style="text-align: center; color: var(--accent-gold); font-style: italic;">... Showing first {max_rows} rows of {len(rows)}. Download full workbook above for complete data. ...</td></tr>')
                    break
                    
            html_parts.append('</table></div>')
        return '\n'.join(html_parts)
    except Exception as e:
        return f"<p class='doc-p'><em>Unable to render Excel preview ({html.escape(str(e))}). Please download the workbook file directly.</em></p>"

def parse_zip_to_html(fpath):
    html_parts = [
        '<div class="archive-list">',
        '  <h3 class="doc-h3">Archive Package Manifest</h3>',
        '  <p class="doc-p" style="margin-bottom: 1rem;">This package contains publication source trees, LaTeX manuscripts, vector graphics, and open-science manifests:</p>',
        '  <div class="table-container"><table class="doc-table">',
        '    <tr><th>File Path</th><th>Uncompressed Size</th></tr>'
    ]
    with zipfile.ZipFile(fpath) as z:
        for item in z.infolist():
            size_kb = max(1, item.file_size // 1024)
            html_parts.append(f'<tr><td style="font-family: var(--font-mono);">{html.escape(item.filename)}</td><td>{size_kb} KB</td></tr>')
            
    html_parts.append('  </table></div></div>')
    return '\n'.join(html_parts)

def parse_code_to_html(fpath, lang="python"):
    with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
        code_text = f.read()
    escaped = html.escape(code_text)
    fname = os.path.basename(fpath)
    return f"""
<div class="code-container">
  <div class="code-header">
    <span>{html.escape(fname)} ({lang.upper()})</span>
    <button class="copy-btn" onclick="navigator.clipboard.writeText(document.getElementById('code-content').innerText); this.innerText='Copied!'; setTimeout(() => this.innerText='Copy Code', 2000);">Copy Code</button>
  </div>
  <pre class="code-block"><code id="code-content">{escaped}</code></pre>
</div>
"""

def parse_markdown_to_html(fpath):
    with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
        md_text = f.read()
    lines = md_text.splitlines()
    html_parts = []
    in_code = False
    code_lines = []
    
    for line in lines:
        if line.strip().startswith('```'):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                escaped_c = html.escape('\n'.join(code_lines))
                html_parts.append(f'<div class="code-container"><pre class="code-block"><code>{escaped_c}</code></pre></div>')
            continue
            
        if in_code:
            code_lines.append(line)
            continue
            
        s_line = line.strip()
        if s_line.startswith('# '):
            html_parts.append(f'<h2 class="doc-h2">{html.escape(s_line[2:])}</h2>')
        elif s_line.startswith('## '):
            html_parts.append(f'<h3 class="doc-h3">{html.escape(s_line[3:])}</h3>')
        elif s_line.startswith('### '):
            html_parts.append(f'<h4 class="doc-h4">{html.escape(s_line[4:])}</h4>')
        elif s_line.startswith('- ') or s_line.startswith('* '):
            html_parts.append(f'<li class="doc-li">{html.escape(s_line[2:])}</li>')
        elif s_line == '---':
            html_parts.append('<hr style="border: none; border-top: 1px solid var(--border-color); margin: 2rem 0;">')
        elif s_line:
            html_parts.append(f'<p class="doc-p">{html.escape(s_line)}</p>')
            
    return '\n'.join(html_parts)

def build_reader_page(paper, prev_p, next_p, base_dir):
    doc_id = paper['id']
    title = paper['title']
    category = paper['category']
    group = paper['group']
    scope = paper['scope']
    status = paper['status']
    file_path = paper.get('file_path', '')
    file_ext = paper.get('file_ext', 'FILE')
    file_size = paper.get('file_size', '')
    associated = paper.get('associated_data', [])
    
    # Generate content HTML
    content_html = ""
    abs_fpath = os.path.join(base_dir, file_path) if file_path else ""
    
    if abs_fpath and os.path.exists(abs_fpath):
        ext_lower = os.path.splitext(abs_fpath)[1].lower()
        if ext_lower == '.docx':
            content_html = parse_docx_to_html(abs_fpath)
        elif ext_lower == '.pdf':
            content_html = f"""
<div class="pdf-viewer-wrap">
  <div class="pdf-toolbar">
    <span>Publication-Grade Proof (Camera-Ready PDF)</span>
    <a href="../{file_path}" target="_blank" class="nav-back-link" style="font-size: 0.78rem;">Open in Full Window &rarr;</a>
  </div>
  <iframe src="../{file_path}" class="pdf-frame"></iframe>
</div>
"""
        elif ext_lower == '.py':
            content_html = parse_code_to_html(abs_fpath, lang="python")
        elif ext_lower == '.json':
            content_html = parse_code_to_html(abs_fpath, lang="json")
        elif ext_lower == '.md':
            content_html = parse_markdown_to_html(abs_fpath)
        elif ext_lower == '.xlsx':
            content_html = parse_xlsx_to_html(abs_fpath)
        elif ext_lower == '.pptx':
            content_html = parse_pptx_to_html(abs_fpath)
        elif ext_lower == '.zip':
            content_html = parse_zip_to_html(abs_fpath)
        elif ext_lower in ['.png', '.jpg', '.jpeg']:
            content_html = f"""
<div class="image-container">
  <img src="../{file_path}" alt="{html.escape(title)}" class="doc-image">
  <p class="image-caption">{html.escape(title)} — {file_size}</p>
</div>
"""
        else:
            content_html = f"<p class='doc-p'>Artifact deliverable: <code>{html.escape(os.path.basename(abs_fpath))}</code></p>"
    else:
        content_html = f"""
<div style="background: rgba(245, 158, 11, 0.08); border: 1px solid rgba(245, 158, 11, 0.25); border-radius: 8px; padding: 1.5rem; margin: 1.5rem 0;">
  <h3 class="doc-h3" style="color: var(--accent-gold); margin-top: 0;">Deliverable File In Preparation</h3>
  <p class="doc-p">This deliverable ({doc_id}) is currently cataloged in the Master Registry. Theoretical scope: {html.escape(scope)}</p>
</div>
"""

    # Build associated data pills HTML
    pills_html = "".join([f'<span class="data-pill">{html.escape(item)}</span>' for item in associated])
    
    # Prev/Next links
    prev_link = f'<a href="{prev_p["id"]}.html" class="nav-arrow-btn">&larr; Prev ({prev_p["id"]})</a>' if prev_p else '<span class="nav-arrow-btn disabled">&larr; Prev</span>'
    next_link = f'<a href="{next_p["id"]}.html" class="nav-arrow-btn">Next ({next_p["id"]}) &rarr;</a>' if next_p else '<span class="nav-arrow-btn disabled">Next &rarr;</span>'
    
    footer_prev = f"""
<a href="{prev_p['id']}.html" class="footer-nav-card">
  <div class="footer-nav-label">&larr; Previous Document</div>
  <div class="footer-nav-title">{prev_p['id']}: {html.escape(prev_p['title'])}</div>
</a>
""" if prev_p else ""

    footer_next = f"""
<a href="{next_p['id']}.html" class="footer-nav-card" style="text-align: right;">
  <div class="footer-nav-label">Next Document &rarr;</div>
  <div class="footer-nav-title">{next_p['id']}: {html.escape(next_p['title'])}</div>
</a>
""" if next_p else ""

    download_btn = f'<a href="../{file_path}" target="_blank" class="nav-download-btn">Download {file_ext} ({file_size})</a>' if file_path else ''

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{doc_id}: {html.escape(title)} — Value Physics Reader</title>
  <meta name="description" content="{html.escape(scope)}">
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Fira+Code:wght@400;600&family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
  <style>
{READER_CSS}
  </style>
</head>
<body>

  <!-- Sticky Top Navigation Bar -->
  <header class="reader-nav">
    <div class="reader-nav-left">
      <a href="../papers.html" class="nav-back-link">&larr; Deliverables Registry</a>
      <span class="doc-badge-id">{doc_id}</span>
      <span class="doc-badge-cat">{html.escape(category)}</span>
      <span class="doc-badge-status">{html.escape(status.split('/')[0].strip())}</span>
    </div>
    <div class="reader-nav-right">
      {prev_link}
      {next_link}
      {download_btn}
    </div>
  </header>

  <!-- Main Article Content Container -->
  <main class="reader-container">
    <article class="article-card">
      <header class="article-header">
        <div class="article-header-meta">
          <span class="data-pill" style="border-color: var(--accent-gold); color: var(--accent-gold);">{html.escape(group)}</span>
          <span style="color: var(--text-dim); font-size: 0.8rem;">Alethekanon Research Institute</span>
        </div>
        <h1 class="article-title">{html.escape(title)}</h1>
        
        <div class="article-author-box">
          <div class="article-author-name">Jarrod Lyndsay Hamilton</div>
          <div class="article-author-inst">Alethekanon Research Institute &middot; Division 2: Societal Economics &amp; Macro-Micro Simulation &middot; Australia</div>
          <div class="article-scope-box">
            <strong>Theoretical &amp; Empirical Scope:</strong> {html.escape(scope)}
          </div>
          {f'<div class="article-data-pills">{pills_html}</div>' if pills_html else ''}
        </div>
      </header>

      <!-- Document Body Content -->
      <div class="doc-content">
        {content_html}
      </div>

      <!-- Footer Navigation Between Papers -->
      <footer class="reader-footer">
        {footer_prev}
        {footer_next}
      </footer>
    </article>
  </main>

  <footer class="site-footer">
    <p>Value Physics Research Repository &middot; Alethekanon Research Institute &middot; CC-BY 4.0 International License</p>
    <p style="margin-top: 0.4rem;"><a href="../index.html">&larr; Return to Master Portal</a> &middot; <a href="../papers.html">All 60 Deliverables</a> &middot; <a href="../data.html">34-Commodity Pricing Matrix</a></p>
  </footer>

</body>
</html>
"""
    return full_html

def generate_all_reader_pages(papers, base_dir):
    read_dir = os.path.join(base_dir, "read")
    os.makedirs(read_dir, exist_ok=True)
    print(f"Generating full-text HTML reader pages in {read_dir}...")
    
    count = 0
    for idx, paper in enumerate(papers):
        prev_p = papers[idx - 1] if idx > 0 else None
        next_p = papers[idx + 1] if idx < len(papers) - 1 else None
        
        page_html = build_reader_page(paper, prev_p, next_p, base_dir)
        out_path = os.path.join(read_dir, f"{paper['id']}.html")
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(page_html)
        count += 1
        
    print(f"Successfully generated {count} reader pages (VPE-01.html to VPE-{count:02d}.html).")

if __name__ == '__main__':
    papers_json_path = os.path.join(BASE_DIR, "data", "papers.json")
    if os.path.exists(papers_json_path):
        with open(papers_json_path, 'r', encoding='utf-8') as f:
            papers = json.load(f)
        generate_all_reader_pages(papers, BASE_DIR)
    else:
        print("Error: papers.json not found.")
